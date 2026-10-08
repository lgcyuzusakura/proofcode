"""Build the editable thesis from reviewed Markdown using bundled python-docx."""
from pathlib import Path
from collections import Counter
import json
import re
import textwrap

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output"
TITLE = "面向可验证代码修改的编程智能体ProofCode的设计与实现"
FIGURES = {
    "architecture": "图3-1 ProofCode总体架构",
    "lifecycle": "图3-2 任务生命周期与批准恢复",
    "retrieval": "图4-1 多路代码检索与证据组织",
    "context": "图5-1 原始会话与模型输入视图",
    "data": "图6-1 数据计划批准与受控执行",
    "experiment": "图7-1 六组对照的执行与比较链路",
}


def fonts(obj, east="宋体", latin="Times New Roman", size=None):
    obj.font.name = latin
    if size is not None:
        obj.font.size = Pt(size)
    obj.font.color.rgb = RGBColor(0, 0, 0)
    if hasattr(obj, "element"):
        rpr = obj.element.get_or_add_rPr()
    else:
        rpr = obj._element.get_or_add_rPr()
    el = rpr.find(qn("w:rFonts"))
    if el is None:
        el = OxmlElement("w:rFonts")
        rpr.insert(0, el)
    for key, value in (("ascii",latin),("hAnsi",latin),("eastAsia",east),("cs",latin)):
        el.set(qn("w:"+key), value)


def field(paragraph, instruction):
    run = paragraph.add_run()
    start=OxmlElement("w:fldChar");start.set(qn("w:fldCharType"),"begin")
    ins=OxmlElement("w:instrText");ins.set(qn("xml:space"),"preserve");ins.text=instruction
    sep=OxmlElement("w:fldChar");sep.set(qn("w:fldCharType"),"separate")
    cached=OxmlElement("w:t");cached.text="1"
    end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end")
    for item in (start,ins,sep,cached,end): run._r.append(item)


def page_number(section, fmt, start=1):
    section.footer.is_linked_to_previous=False
    p=section.footer.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    field(p," PAGE ")
    for r in p.runs:fonts(r,size=10.5)
    num=OxmlElement("w:pgNumType");num.set(qn("w:fmt"),fmt);num.set(qn("w:start"),str(start))
    section._sectPr.append(num)


def base_section(section):
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(2.5);section.bottom_margin=Cm(2.5)
    section.left_margin=Cm(2.7);section.right_margin=Cm(2.5)
    section.header_distance=Cm(1.2);section.footer_distance=Cm(1.2)


def styles(doc):
    # The bundled template contains a blue bottom border on Title.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal=doc.styles["Normal"]
    fonts(normal,size=12)
    pf=normal.paragraph_format
    pf.line_spacing=Pt(20);pf.space_after=Pt(0);pf.space_before=Pt(0)
    pf.first_line_indent=Pt(24);pf.widow_control=True
    for name,size in (("Title",24),("Heading 1",16),("Heading 2",14),("Heading 3",12)):
        s=doc.styles[name];fonts(s,"黑体",size=size)
        s.font.bold=True
        p=s.paragraph_format
        p.first_line_indent=Pt(0);p.keep_with_next=True;p.keep_together=True
        p.space_before=Pt(12 if name != "Heading 1" else 0);p.space_after=Pt(8)
        p.line_spacing=Pt(24 if name in ("Heading 1","Title") else 22)
    doc.styles["Heading 1"].paragraph_format.page_break_before=True
    doc.styles["Title"].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if "TOC Heading" not in doc.styles:
        doc.styles.add_style("TOC Heading",1)
    toc_heading=doc.styles["TOC Heading"]
    toc_heading.base_style=doc.styles["Heading 1"]
    fonts(toc_heading,"黑体",size=16)
    toc_heading.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    outline=OxmlElement("w:outlineLvl");outline.set(qn("w:val"),"9")
    toc_heading.element.get_or_add_pPr().append(outline)
    caption=doc.styles["Caption"];fonts(caption,size=10.5)
    caption.paragraph_format.first_line_indent=Pt(0)
    caption.paragraph_format.line_spacing=Pt(16)
    caption.paragraph_format.space_after=Pt(6)
    caption.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for name in ["TOC 1","TOC 2","TOC 3"]:
        if name not in doc.styles:
            doc.styles.add_style(name,1)
        s=doc.styles[name];fonts(s,size=12)
        s.paragraph_format.first_line_indent=Pt(0)
        s.paragraph_format.line_spacing=Pt(18)
        s.paragraph_format.space_after=Pt(0)
    if "Code Block" not in doc.styles:doc.styles.add_style("Code Block",1)
    code=doc.styles["Code Block"];fonts(code,latin="Consolas",size=9.5)
    code.paragraph_format.first_line_indent=Pt(0)
    code.paragraph_format.left_indent=Cm(.3);code.paragraph_format.right_indent=Cm(.2)
    code.paragraph_format.line_spacing=Pt(13);code.paragraph_format.space_after=Pt(6)
    code.paragraph_format.keep_together=True


def inline(paragraph, token):
    bold=False;italic=False
    for child in token.children or []:
        if child.type == "strong_open":bold=True;continue
        if child.type == "strong_close":bold=False;continue
        if child.type == "em_open":italic=True;continue
        if child.type == "em_close":italic=False;continue
        if child.type in ("link_open","link_close"):continue
        if child.type in ("softbreak","hardbreak"):
            paragraph.add_run().add_break();continue
        if child.type in ("text","code_inline"):
            r=paragraph.add_run(child.content);r.bold=bold;r.italic=italic
            if child.type == "code_inline":fonts(r,latin="Consolas",size=10.5)


def table(doc, tokens):
    rows=[];current=None
    for i,t in enumerate(tokens):
        if t.type=="tr_open":current=[]
        elif t.type=="inline" and current is not None:current.append(t)
        elif t.type=="tr_close":rows.append(current);current=None
    if not rows:return
    count=len(rows[0]);tbl=doc.add_table(rows=0,cols=count)
    tbl.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl.autofit=False
    available=15.8
    ratios={2:[.25,.75],3:[.25,.4,.35],4:[.18,.27,.28,.27],5:[.10,.23,.13,.28,.26]}.get(count,[1/count]*count)
    headers=[token.content for token in rows[0]]
    if headers == ["文件","记录内容","可以支持的结论"]:
        ratios=[.43,.28,.29]
    elif headers == ["项目","文献编号","本课题参考内容"]:
        ratios=[.26,.16,.58]
    for col,ratio in zip(tbl.columns,ratios):col.width=Cm(available*ratio)
    props=tbl._tbl.tblPr
    borders=OxmlElement("w:tblBorders")
    for kind in ["top","bottom","left","right","insideH","insideV"]:
        edge=OxmlElement("w:"+kind);edge.set(qn("w:val"),"single")
        edge.set(qn("w:sz"),"4");edge.set(qn("w:color"),"D9D9D9");borders.append(edge)
    props.append(borders)
    margins=OxmlElement("w:tblCellMar")
    for edge,size in [("top",60),("bottom",60),("left",70),("right",70)]:
        el=OxmlElement("w:"+edge);el.set(qn("w:w"),str(size));el.set(qn("w:type"),"dxa");margins.append(el)
    props.append(margins)
    for index,row in enumerate(rows):
        cells=tbl.add_row().cells
        pr=tbl.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement("w:cantSplit"))
        if index==0:pr.append(OxmlElement("w:tblHeader"))
        for j,t in enumerate(row):
            cell=cells[j];cell.width=Cm(available*ratios[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=Pt(16);p.paragraph_format.keep_with_next=False
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(t.content)<=14 else WD_ALIGN_PARAGRAPH.LEFT
            inline(p,t)
            for run in p.runs:fonts(run,size=10.5);run.bold=index==0
            if index==0:
                shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"EEEEEE")
                cell._tc.get_or_add_tcPr().append(shade)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.line_spacing=Pt(4);p.paragraph_format.space_after=Pt(4)


def figure(doc,key):
    assert key in FIGURES,key
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(6)
    p.paragraph_format.line_spacing=1.0
    p.add_run().add_picture(str(ROOT / "figures" / (key+".png")),width=Cm(15.5))
    cap=doc.add_paragraph(FIGURES[key],style="Caption")
    cap.paragraph_format.keep_with_next=False


def render_md(doc, content, first=False, abstract=False, references=False):
    parser=MarkdownIt("commonmark").enable("table")
    tokens=parser.parse(content)
    i=0;first_heading=first
    while i<len(tokens):
        t=tokens[i]
        if t.type=="heading_open":
            level=int(t.tag[1:]);p=doc.add_paragraph(style="Heading "+str(level));inline(p,tokens[i+1])
            if level==1:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                if first_heading:p.paragraph_format.page_break_before=False;first_heading=False
            i+=3;continue
        if t.type=="paragraph_open":
            text=tokens[i+1].content
            match=re.fullmatch(r"!FIG\[(\w+)\]",text)
            if match:figure(doc,match[1]);i+=3;continue
            p=doc.add_paragraph();inline(p,tokens[i+1])
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            if re.match(r"^表\d+-\d+",text):
                p.style="Caption";p.paragraph_format.keep_with_next=True
            elif text.startswith(("关键词：","Keywords:")):
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(8)
            elif references:
                p.paragraph_format.first_line_indent=Cm(-.7);p.paragraph_format.left_indent=Cm(.7)
                p.paragraph_format.line_spacing=Pt(16);p.paragraph_format.space_after=Pt(6)
                for run in p.runs:fonts(run,size=10.5)
            elif abstract and not re.search("[\u4e00-\u9fff]",text):
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(17)
                p.paragraph_format.space_after=Pt(6)
            i+=3;continue
        if t.type=="fence":
            p=doc.add_paragraph(style="Code Block")
            lines=[]
            for line in t.content.rstrip().splitlines():
                if len(line)>76:
                    lines.extend(textwrap.wrap(line,width=76,break_long_words=True,break_on_hyphens=False,replace_whitespace=False,drop_whitespace=False) or [""])
                else:lines.append(line)
            p.add_run("\n".join(lines))
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"F5F5F5");p._p.get_or_add_pPr().append(shade)
            i+=1;continue
        if t.type=="table_open":
            end=i+1
            while tokens[end].type!="table_close":end+=1
            table(doc,tokens[i:end+1]);i=end+1;continue
        i+=1


def cover(doc):
    p=doc.add_paragraph("长春大学",style="Title");p.paragraph_format.space_before=Pt(45)
    p=doc.add_paragraph("本科毕业设计（论文）")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_before=Pt(25);p.paragraph_format.space_after=Pt(48)
    for r in p.runs:fonts(r,"黑体",size=20)
    p=doc.add_paragraph("面向可验证代码修改的编程智能体\nProofCode的设计与实现",style="Title")
    p.paragraph_format.line_spacing=Pt(32);p.paragraph_format.space_after=Pt(55)
    for r in p.runs:fonts(r,"黑体",size=20)
    for label,value in [("学    院","________________________"),("专    业","计算机科学与技术"),("学生姓名","________________________"),("学    号","________________________"),("指导教师","________________________"),("毕业届别","________________________")]:
        p=doc.add_paragraph(label+"    "+value)
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Cm(2.4)
        p.paragraph_format.line_spacing=Pt(28)
        for r in p.runs:fonts(r,size=14)
    p=doc.add_paragraph("年    月")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_before=Pt(45)


def build():
    OUTPUT.mkdir(parents=True,exist_ok=True)
    doc=Document();styles(doc);base_section(doc.sections[0])
    doc.core_properties.title=TITLE
    doc.core_properties.author=""
    doc.core_properties.subject="计算机科学与技术本科毕业设计论文"
    cover(doc)
    pre=doc.add_section(WD_SECTION_START.NEW_PAGE);base_section(pre);page_number(pre,"upperRoman")
    abstract=(ROOT/"manuscript/00-abstract.md").read_text(encoding="utf-8")
    render_md(doc,abstract,first=True,abstract=True)
    p=doc.add_paragraph("目录",style="TOC Heading")
    p=doc.add_paragraph("TOC_INSERT_POINT");p.paragraph_format.first_line_indent=Pt(0)
    body=doc.add_section(WD_SECTION_START.NEW_PAGE);base_section(body);page_number(body,"decimal")
    for n,path in enumerate(sorted((ROOT/"manuscript").glob("*.md"))):
        if path.name.startswith("00-"):continue
        render_md(doc,path.read_text(encoding="utf-8"),first=path.name.startswith("01-"),references=path.name.startswith("10-"))
    settings=doc.settings.element
    update=OxmlElement("w:updateFields");update.set(qn("w:val"),"true");settings.append(update)
    output=OUTPUT/"ProofCode_本科毕业设计论文_初稿.docx"
    doc.save(output)
    body_text="\n".join(p.read_text(encoding="utf-8") for p in sorted((ROOT/"manuscript").glob("0[1-8]-*.md")))
    all_text="\n\n".join(p.read_text(encoding="utf-8") for p in sorted((ROOT/"manuscript").glob("*.md")))
    counts={"bodyChineseCharacters":len(re.findall(r"[\u4e00-\u9fff]",body_text)),"allChineseCharacters":len(re.findall(r"[\u4e00-\u9fff]",all_text)),"bodyCharactersWithoutWhitespace":len(re.sub(r"\s","",body_text)),"wordCountIsProvisional":True,"figures":len(FIGURES),"references":35}
    (ROOT/"evidence/manuscript-counts.json").write_text(json.dumps(counts,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    (OUTPUT/"ProofCode_论文正文.md").write_text("# "+TITLE+"\n\n"+all_text.rstrip()+"\n",encoding="utf-8")
    print(json.dumps({"output":str(output),**counts},ensure_ascii=False))


if __name__=="__main__":
    build()
